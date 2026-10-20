import io
import unicodedata
import pandas as pd
from django.db.models import QuerySet
from analisis.models import ResultadoTaxonomico, CatalogoMicroorganismo

TAXONOMIC_STOPWORDS = {
    'unclassified', 'unknown', 'uncultured', 'unidentified', 'sp', 'nan', 'bacterium',
    'unidentified bacterium', 'unidentified eubacterium', 'uncultured soil'
}

class CatalogService:
    
    @staticmethod
    def _limpiar(val):
        return str(val).strip().lower() if val is not None else ""

    @classmethod
    def obtener_catalogo_procesado(cls, queryset_resultados: QuerySet, clasificacion="", nivel=""):
        """
        Procesa el matching jerárquico recibiendo el QuerySet pre-filtrado de la vista.
        Resuelve definitivamente el error de tipos en Pylance.
        """
        # 1. Carga maestra para evitar N+1
        reglas = CatalogoMicroorganismo.objects.all()
        # Mapa: {(nivel, nombre): clasificacion}
        dicc_maestro = {(r.nivel.lower(), r.nombre.lower()): r.clasificacion for r in reglas}

        resultados = []
        stats = {"total": 0, "benefico": 0, "fitopatogeno": 0, "neutral": 0, "sin_clasificar": 0}

        # Procesamos el QuerySet que llega desde la vista
        for r in queryset_resultados:
            genus_limpio = cls._limpiar(r.genus)
            species_limpio = cls._limpiar(r.species)

            nombre = genus_limpio
            nivel_tax = "genus"
            tipo = None

            # Matching jerárquico adaptativo
            if species_limpio and species_limpio not in TAXONOMIC_STOPWORDS:
                if species_limpio.startswith(genus_limpio):
                    binomio = species_limpio
                else:
                    binomio = f"{genus_limpio} {species_limpio}"

                # 1. Intentar matching de especie
                tipo = dicc_maestro.get(("species", binomio))
                if tipo:
                    nivel_tax = "species"
                    nombre = binomio
                else:
                    # 2. Fallback a género
                    tipo = dicc_maestro.get(("genus", genus_limpio))
                    nivel_tax = "genus"
                    nombre = genus_limpio
            else:
                # Fallback directo a género
                tipo = dicc_maestro.get(("genus", genus_limpio))
                nivel_tax = "genus"
                nombre = genus_limpio

            # Si no hubo match en el catálogo maestro, asegurar "sin_clasificar"
            if not tipo:
                tipo = "sin_clasificar"

            # Aplicar filtros dinámicos en memoria (requerido porque 'tipo' no existe en la BD)
            if (clasificacion and clasificacion != tipo) or (nivel and nivel != nivel_tax):
                continue

            # Incremento de estadísticas tras clasificación final
            stats["total"] += 1
            stats[tipo] = stats.get(tipo, 0) + 1

            resultados.append({
                "id": r.id,
                "taxonomia": nombre,
                "nivel": nivel_tax,
                "producto": {
                    "tipo": r.archivo.tipo_analisis if hasattr(r.archivo, 'tipo_analisis') else "",
                    "nombre": r.archivo.muestra.codigo if hasattr(r.archivo, 'muestra') and hasattr(r.archivo.muestra, 'codigo') else "Muestra S/N"
                },
                "porcentaje_hits": r.porcentaje_hits,
                "clasificacion": tipo
            })

        # Adaptación de nombres de contadores para tu plantilla HTML
        estadisticas = {
            "total": stats["total"],
            "total_beneficos": stats["benefico"],
            "total_fitopatogenos": stats["fitopatogeno"],
            "total_neutrales": stats["neutral"],
            "total_sin_clas": stats["sin_clasificar"]
        }
        return resultados, estadisticas

    @staticmethod
    def _quitar_tildes(texto: str) -> str:
        nfkd = unicodedata.normalize('NFD', texto)
        return ''.join(c for c in nfkd if unicodedata.category(c) != 'Mn')

    @classmethod
    def _normalizar_nivel(cls, valor: str) -> str:
        """
        Acepta variantes en español e inglés (con/sin tilde, mayúsculas, abreviaturas)
        y retorna el valor canónico: genus | species | family | order | class | phylum.
        Si no reconoce el valor devuelve 'genus' como fallback seguro.
        """
        clave = cls._quitar_tildes(str(valor).strip().lower())
        MAPA = {
            # genus
            'genus': 'genus', 'genero': 'genus', 'gen': 'genus', 'g': 'genus',
            # species
            'species': 'species', 'especie': 'species', 'specie': 'species',
            'sp': 'species', 'spp': 'species',
            # family
            'family': 'family', 'familia': 'family', 'fam': 'family',
            # order
            'order': 'order', 'orden': 'order', 'ord': 'order',
            # class
            'class': 'class', 'clase': 'class', 'cls': 'class',
            # phylum
            'phylum': 'phylum', 'filo': 'phylum', 'division': 'phylum',
        }
        return MAPA.get(clave, 'genus')

    @classmethod
    def _normalizar_clasificacion(cls, valor: str) -> str:
        """
        Acepta cualquier variante legible del Excel y la mapea al código interno.
        Maneja acentos, mayúsculas, sinónimos en español e inglés.
        """
        v = cls._quitar_tildes(str(valor).strip().lower())
        MAPA = {
            # benéfico
            'benefico':            'benefico',
            'beneficioso':         'benefico',
            'beneficial':          'benefico',
            'promotor':            'benefico',
            'aliado':              'benefico',
            # fitopatógeno
            'fitopatogeno':        'fitopatogeno',
            'potencial fitopatogeno': 'fitopatogeno',
            'patogeno':            'fitopatogeno',
            'patogeno potencial':  'fitopatogeno',
            'pathogen':            'fitopatogeno',
            'riesgo':              'fitopatogeno',
            # neutral
            'neutral':             'neutral',
            'neutro':              'neutral',
            # sin clasificar
            'sin_clasificar':      'sin_clasificar',
            'sin clasificar':      'sin_clasificar',
            'desconocido':         'sin_clasificar',
            'unknown':             'sin_clasificar',
            'nd':                  'sin_clasificar',
            'n/d':                 'sin_clasificar',
            '':                    'sin_clasificar',
            'nan':                 'sin_clasificar',
            'none':                'sin_clasificar',
        }
        return MAPA.get(v, 'sin_clasificar')

    @staticmethod
    def importar_desde_archivo(archivo) -> tuple:
        """
        Lee un CSV/Excel y crea o actualiza registros en CatalogoMicroorganismo.
        Columnas requeridas: nombre, nivel, marcador (16S/ITS), clasificacion.
        Columnas opcionales: descripcion, referencia.
        Retorna (creados, actualizados, errores).
        """
        nombre_archivo = archivo.name.lower()
        if nombre_archivo.endswith(('.xlsx', '.xls')):
            df = pd.read_excel(archivo)
        else:
            archivo.seek(0)
            raw = archivo.read()
            try:
                df = pd.read_csv(io.BytesIO(raw))
            except Exception:
                df = pd.read_csv(io.BytesIO(raw), sep='\t')

        df.columns = [str(c).strip().lower() for c in df.columns]

        required = {'nombre', 'nivel', 'marcador', 'clasificacion'}
        missing = required - set(df.columns)
        if missing:
            raise ValueError(
                f"Columnas requeridas faltantes: {', '.join(sorted(missing))}. "
                "El archivo debe tener: nombre, nivel, marcador, clasificacion "
                "(opcionales: descripcion, referencia)"
            )

        MARCADORES = {'16S', 'ITS'}
        _nan = {'nan', 'none', ''}

        creados = actualizados = errores = 0

        for _, row in df.iterrows():
            nombre = CatalogService._quitar_escape_formula(str(row.get('nombre', '')).strip())
            nivel  = CatalogService._normalizar_nivel(row.get('nivel', 'genus'))

            if not nombre or nombre.lower() in _nan:
                errores += 1
                continue

            marcador = str(row.get('marcador', '16S')).strip().upper()
            if marcador not in MARCADORES:
                marcador = '16S'

            # Normalización robusta: acepta tildes y sinónimos del Excel
            clasificacion = CatalogService._normalizar_clasificacion(
                row.get('clasificacion', 'sin_clasificar')
            )

            _nan_set = _nan
            def _clean(val):
                s = CatalogService._quitar_escape_formula(str(val).strip())
                return '' if s.lower() in _nan_set else s

            descripcion = _clean(row.get('descripcion', '')) if 'descripcion' in df.columns else ''
            referencia  = _clean(row.get('referencia',  '')) if 'referencia'  in df.columns else ''

            try:
                # Django 5.x: update_or_create no acepta __iexact en kwargs no-defaults.
                # Usamos filter().first() para la búsqueda case-insensitive y luego
                # create/save explícito para la escritura.
                obj = CatalogoMicroorganismo.objects.filter(
                    nombre__iexact=nombre,
                    nivel__iexact=nivel,
                ).first()

                if obj:
                    obj.marcador      = marcador
                    obj.clasificacion = clasificacion
                    obj.descripcion   = descripcion or obj.descripcion
                    obj.referencia    = referencia  or obj.referencia
                    obj.save()
                    actualizados += 1
                else:
                    CatalogoMicroorganismo.objects.create(
                        nombre=nombre,
                        nivel=nivel,
                        marcador=marcador,
                        clasificacion=clasificacion,
                        descripcion=descripcion,
                        referencia=referencia,
                    )
                    creados += 1
            except Exception:
                errores += 1

        return creados, actualizados, errores

    # ---------------------------
    # EXPORTACIÓN
    # ---------------------------

    # Mismas columnas (y mismo orden) que exige importar_desde_archivo, para
    # que un archivo exportado se pueda volver a importar tal cual.
    COLUMNAS_EXPORTACION = ['nombre', 'nivel', 'marcador', 'clasificacion', 'descripcion', 'referencia']

    @staticmethod
    def _quitar_escape_formula(texto: str) -> str:
        """
        Revierte el apóstrofe que la exportación antepone a textos que Excel
        interpretaría como fórmula ('=…, '+…, '-…, '@…), para que exportar y
        volver a importar no altere los datos.
        """
        from analisis.services.business.export_service import _EXCEL_FORMULA_TRIGGERS
        if texto.startswith("'") and texto[1:2] and texto[1:].startswith(_EXCEL_FORMULA_TRIGGERS):
            return texto[1:]
        return texto

    @staticmethod
    def _dataframe_exportacion(filas) -> pd.DataFrame:
        """filas: iterable de dicts con las COLUMNAS_EXPORTACION (p. ej. queryset.values(...))."""
        from analisis.services.business.export_service import _excel_safe_df
        df = pd.DataFrame(list(filas), columns=CatalogService.COLUMNAS_EXPORTACION).fillna('')
        return _excel_safe_df(df)

    @staticmethod
    def exportar_csv(filas) -> bytes:
        """CSV UTF-8 con BOM (para que Excel muestre bien las tildes), reimportable."""
        df = CatalogService._dataframe_exportacion(filas)
        return df.to_csv(index=False).encode('utf-8-sig')

    @staticmethod
    def exportar_excel(filas) -> bytes:
        """
        Libro con una hoja «Catálogo» (todas las filas; es la que lee la
        importación, que toma la primera hoja) y una hoja por marcador.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        df = CatalogService._dataframe_exportacion(filas)
        verde = PatternFill('solid', fgColor='0F6E56')
        insignias = {
            'benefico':       ('D4EDDA', '14532D'),
            'fitopatogeno':   ('FADBD8', '7F1D1D'),
            'neutral':        ('D6E4F0', '1E3A5F'),
            'sin_clasificar': ('E5E7EB', '374151'),
        }
        anchos = {'nombre': 38, 'nivel': 10, 'marcador': 10, 'clasificacion': 16,
                  'descripcion': 60, 'referencia': 50}
        col_clasif = CatalogService.COLUMNAS_EXPORTACION.index('clasificacion') + 1

        wb = Workbook()
        wb.remove(wb.active)
        hojas = [('Catálogo', df)] + [
            (titulo, df[df['marcador'] == marcador])
            for marcador, titulo in (('16S', 'Bacterias (16S)'), ('ITS', 'Hongos (ITS)'))
        ]
        for titulo, datos in hojas:
            if titulo != 'Catálogo' and datos.empty:
                continue
            ws = wb.create_sheet(titulo)
            ws.append(CatalogService.COLUMNAS_EXPORTACION)
            for fila in datos.itertuples(index=False):
                ws.append(list(fila))

            for j, col in enumerate(CatalogService.COLUMNAS_EXPORTACION, start=1):
                celda = ws.cell(row=1, column=j)
                celda.fill = verde
                celda.font = Font(bold=True, color='FFFFFF')
                celda.alignment = Alignment(horizontal='center', vertical='center')
                ws.column_dimensions[get_column_letter(j)].width = anchos[col]
            for (celda,) in ws.iter_rows(min_row=2, min_col=col_clasif, max_col=col_clasif):
                bg, fg = insignias.get(celda.value, insignias['sin_clasificar'])
                celda.fill = PatternFill('solid', fgColor=bg)
                celda.font = Font(bold=True, color=fg)
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions

        salida = io.BytesIO()
        wb.save(salida)
        return salida.getvalue()

    @staticmethod
    def reclasificar_microorganismo(pk, nueva_clasificacion, descripcion):
        try:
            res = ResultadoTaxonomico.objects.select_related('archivo').get(pk=pk)
            marcador = res.archivo.tipo_analisis  # "16S" o "ITS"

            genus_limpio = CatalogService._limpiar(res.genus)
            species_limpio = CatalogService._limpiar(res.species)

            if species_limpio and species_limpio not in TAXONOMIC_STOPWORDS:
                nombre = f"{genus_limpio} {species_limpio}".strip()
                nivel = "species"
            else:
                nombre = genus_limpio
                nivel = "genus"

            obj = CatalogoMicroorganismo.objects.filter(
                nombre__iexact=nombre,
                nivel__iexact=nivel,
                marcador=marcador,
            ).first()
            if obj:
                obj.clasificacion = nueva_clasificacion
                obj.descripcion   = descripcion or obj.descripcion
                obj.save()
            else:
                CatalogoMicroorganismo.objects.create(
                    nombre=nombre,
                    nivel=nivel,
                    marcador=marcador,
                    clasificacion=nueva_clasificacion,
                    descripcion=descripcion,
                )
            return True
        except:
            return False